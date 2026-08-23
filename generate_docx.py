from pathlib import Path
from docx import Document

out = Path('frep-explanation.docx')
doc = Document()
doc.add_heading('Factory Resource Exchange Platform (FREP)', level=1)
doc.add_paragraph('FREP turns idle industrial assets into a shared production network that helps MSMEs reduce cost, increase uptime, and collaborate smarter.').style = 'Intense Quote'
doc.add_heading('Problem', level=2)
doc.add_paragraph('Many MSMEs cannot afford expensive machinery, testing equipment, skilled manpower, or logistics support at all times. At the same time, other factories own these same resources but leave them idle during low-demand periods.')
doc.add_heading('Idea', level=2)
doc.add_paragraph('FREP is a digital platform that connects verified MSMEs so they can share industrial resources instead of buying and owning everything individually.')
doc.add_heading('What makes it unique', level=2)
doc.add_paragraph('The most distinctive feature of FREP is its AI-style matching logic. It evaluates category fit, location proximity, budget fit, availability, and verification bonus.')
doc.add_heading('Objective', level=2)
doc.add_paragraph('The objective of FREP is to improve the utilization of industrial assets across the MSME ecosystem, reduce capital expenditure, and create a digital culture of collaborative manufacturing.')
doc.add_heading('How the platform works', level=2)
doc.add_paragraph('A factory lists a resource, it is verified, a business submits a need, the AI matching engine scores resources, a booking is placed, and feedback is shared after use.')
doc.add_heading('Where it applies', level=2)
doc.add_paragraph('FREP can be used in textile, automotive, metal fabrication, plastics, packaging, electronics, food processing, pharma, industrial parks, and MSME clusters.')
doc.add_heading('Market potential', level=2)
doc.add_paragraph('India has a very large MSME base, and many of these units face cost and access barriers in everyday operations. FREP creates a commercial ecosystem through booking commissions, subscriptions, logistics partnership revenue, and AI-enabled services.')
rows = [
    ['Aspect', 'FREP Value'],
    ['Core problem', 'Idle industrial resources and limited access for MSMEs'],
    ['Main users', 'Buyers, resource owners, and network admins'],
    ['Core advantage', 'Explainable AI matching for better resource selection'],
    ['Business model', 'Subscriptions, booking commissions, logistics partnerships, analytics services'],
]
table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
header_cells = table.rows[0].cells
header_cells[0].text = rows[0][0]
header_cells[1].text = rows[0][1]
for row in rows[1:]:
    cells = table.add_row().cells
    cells[0].text = row[0]
    cells[1].text = row[1]

doc.save(out)
print('saved', out.exists(), out.stat().st_size)
